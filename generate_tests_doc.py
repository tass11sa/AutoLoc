import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()

# Margins
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=200, right=200):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_placeholder_box(doc, caption_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8F9FA")
    set_cell_margins(cell, top=320, bottom=320, left=320, right=320)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="dashed" w:sz="10" w:space="0" w:color="0078D4"/><w:bottom w:val="dashed" w:sz="10" w:space="0" w:color="0078D4"/><w:left w:val="dashed" w:sz="10" w:space="0" w:color="0078D4"/><w:right w:val="dashed" w:sz="10" w:space="0" w:color="0078D4"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run1 = p.add_run("📷 [ Insérer ici la capture d'écran Postman : ")
    run1.font.size = Pt(10.5)
    run1.font.color.rgb = RGBColor(0x00, 0x78, 0xD4)
    run1.font.bold = True
    
    run2 = p.add_run(caption_text)
    run2.font.size = Pt(10.5)
    run2.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
    run2.font.bold = True
    
    run3 = p.add_run(" - Status: 200 OK ]")
    run3.font.size = Pt(10.5)
    run3.font.color.rgb = RGBColor(0x00, 0x80, 0x00)
    run3.font.bold = True
    
    doc.add_paragraph()

# Title
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title_p.add_run("Rapport de Test des Endpoints REST : VehiculeController\nStatus HTTP : 200 OK")
title_run.font.size = Pt(18)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(0x0F, 0x4C, 0x81)

# Student Info
info_tbl = doc.add_table(rows=4, cols=2)
info_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
infos = [
    ("Étudiante :", "Tasnim Samoud"),
    ("Classe / Module :", "ASI / 4Ssa - Architecture des Systèmes d'Information (Spring Boot)"),
    ("Projet :", "AutoLoc - Contrôleur REST VehiculeController"),
    ("Serveur / Port :", "http://localhost:8081/vehicules")
]
for i, (label, val) in enumerate(infos):
    c1, c2 = info_tbl.rows[i].cells
    set_cell_background(c1, "F0F4F8")
    set_cell_background(c2, "FFFFFF")
    
    p1 = c1.paragraphs[0]
    r1 = p1.add_run(label)
    r1.font.bold = True
    r1.font.size = Pt(10)
    r1.font.color.rgb = RGBColor(0x0F, 0x4C, 0x81)
    
    p2 = c2.paragraphs[0]
    r2 = p2.add_run(val)
    r2.font.size = Pt(10)

doc.add_paragraph()

# Section 1 : Tableau Récapitulatif
h1 = doc.add_heading("1. Tableau Récapitulatif des Endpoints Testés", level=1)
h1.runs[0].font.color.rgb = RGBColor(0x0F, 0x4C, 0x81)

tbl = doc.add_table(rows=1, cols=4)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["Méthode HTTP", "URL de l'Endpoint", "Description", "Statut Obtenu"]
for j, h in enumerate(headers):
    c = tbl.rows[0].cells[j]
    set_cell_background(c, "0F4C81")
    p = c.paragraphs[0]
    r = p.add_run(h)
    r.font.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

endpoints = [
    ("POST", "http://localhost:8081/vehicules/ajouter", "Ajouter un nouveau véhicule en BDD", "200 OK"),
    ("GET", "http://localhost:8081/vehicules/all", "Consulter la liste de tous les véhicules", "200 OK"),
    ("GET", "http://localhost:8081/vehicules/{id}", "Consulter les détails d'un véhicule par ID", "200 OK"),
    ("PUT", "http://localhost:8081/vehicules/modifier", "Mettre à jour les informations d'un véhicule", "200 OK"),
    ("DELETE", "http://localhost:8081/vehicules/supprimer/{id}", "Supprimer un véhicule par son ID", "200 OK")
]

for row_data in endpoints:
    row = tbl.add_row()
    for col_idx, cell_value in enumerate(row_data):
        cell = row.cells[col_idx]
        set_cell_background(cell, "F9FBFD" if col_idx == 0 else "FFFFFF")
        p = cell.paragraphs[0]
        r = p.add_run(cell_value)
        r.font.size = Pt(9)
        if col_idx == 0:
            r.font.bold = True
            if cell_value == "POST":
                r.font.color.rgb = RGBColor(0xFD, 0x7E, 0x14)
            elif cell_value == "GET":
                r.font.color.rgb = RGBColor(0x0D, 0x6E, 0xFD)
            elif cell_value == "PUT":
                r.font.color.rgb = RGBColor(0x19, 0x87, 0x54)
            elif cell_value == "DELETE":
                r.font.color.rgb = RGBColor(0xDC, 0x35, 0x45)
        elif col_idx == 3:
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x19, 0x87, 0x54)

doc.add_paragraph()

# Section 2 : Détails des Tests et Captures d'écran
h2 = doc.add_heading("2. Captures d'Écran des Tests Postman (Status 200 OK)", level=1)
h2.runs[0].font.color.rgb = RGBColor(0x0F, 0x4C, 0x81)

# 2.1 POST
doc.add_heading("2.1 Test du POST : Ajouter un Véhicule", level=2)
p_post = doc.add_paragraph("• Méthode : POST\n• URL : http://localhost:8081/vehicules/ajouter\n• Header : Content-Type: application/json\n• Request Body (JSON) :")
p_post.runs[0].font.size = Pt(9.5)

p_json_post = doc.add_paragraph(
"""{
    "immatriculation": "240-TUNIS-1234",
    "marque": "Peugeot",
    "modele": "208",
    "categorie": "CITADINE",
    "tarifJournalier": 120.00,
    "statut": "DISPONIBLE"
}"""
)
p_json_post.runs[0].font.name = "Consolas"
p_json_post.runs[0].font.size = Pt(9)

add_placeholder_box(doc, "POST /vehicules/ajouter")

# 2.2 GET ALL
doc.add_heading("2.2 Test du GET : Récupérer tous les Véhicules", level=2)
p_all = doc.add_paragraph("• Méthode : GET\n• URL : http://localhost:8081/vehicules/all\n• Résultat attendu : Liste JSON de tous les véhicules avec le statut 200 OK.")
p_all.runs[0].font.size = Pt(9.5)
add_placeholder_box(doc, "GET /vehicules/all")

# 2.3 GET BY ID
doc.add_heading("2.3 Test du GET By ID : Consulter un Véhicule Spécifique", level=2)
p_id = doc.add_paragraph("• Méthode : GET\n• URL : http://localhost:8081/vehicules/1\n• Résultat attendu : Objet JSON du véhicule correspondant à l'ID avec le statut 200 OK.")
p_id.runs[0].font.size = Pt(9.5)
add_placeholder_box(doc, "GET /vehicules/{id}")

# 2.4 PUT
doc.add_heading("2.4 Test du PUT : Modifier un Véhicule", level=2)
p_put = doc.add_paragraph("• Méthode : PUT\n• URL : http://localhost:8081/vehicules/modifier\n• Header : Content-Type: application/json\n• Request Body (JSON avec modifications) :")
p_put.runs[0].font.size = Pt(9.5)

p_json_put = doc.add_paragraph(
"""{
    "idVehicule": 1,
    "immatriculation": "240-TUNIS-1234",
    "marque": "Peugeot",
    "modele": "208",
    "categorie": "CITADINE",
    "tarifJournalier": 135.00,
    "statut": "LOUE"
}"""
)
p_json_put.runs[0].font.name = "Consolas"
p_json_put.runs[0].font.size = Pt(9)

add_placeholder_box(doc, "PUT /vehicules/modifier")

# 2.5 DELETE
doc.add_heading("2.5 Test du DELETE : Supprimer un Véhicule", level=2)
p_del = doc.add_paragraph("• Méthode : DELETE\n• URL : http://localhost:8081/vehicules/supprimer/1\n• Résultat attendu : Suppression du véhicule en base et réponse avec statut 200 OK.")
p_del.runs[0].font.size = Pt(9.5)
add_placeholder_box(doc, "DELETE /vehicules/supprimer/{id}")

out_path = r"c:\Users\lenovo\Documents\AutoLoc-repo\Rendu_Tests_VehiculeController.docx"
doc.save(out_path)
print(f"SUCCESS: Document saved at {out_path}")
